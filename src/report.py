# -*- coding: utf-8 -*-
import pandas as pd
from config import settings as S
def _style_sheet(ws, header_fill='1F4E78', header_font_color='FFFFFF'):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    fill = PatternFill('solid', fgColor=header_fill)
    font = Font(bold=True, color=header_font_color, size=11)
    thin = Side(style='thin', color='D9D9D9')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        mx = max((len(str(c.value)) if c.value is not None else 0 for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(mx + 4, 45)
def export_excel(kpis, monthly, top, cat, city, pay, status, out=None, profit=None, pprod=None, pcat=None,
                 basket=None, cohort=None, forecast_info=None, forecast_series=None, rfm=None, seg=None,
                 restock=None, sentiment_summary=None, sentiment_byprod=None, alerts=None, recs=None, repeat=None,
                 quality=None, churn=None, churn_summary=None, tiers=None, tier_corr=None, gaps=None,
                 recoverable=None, upsell=None):
    out = out or S.EXCEL_REPORT_FILE
    S.DATA_REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out, engine='openpyxl') as xw:
        pd.DataFrame([kpis]).to_excel(xw, sheet_name='KPIs', index=False)
        monthly.to_excel(xw, sheet_name='Monthly', index=False)
        top.to_excel(xw, sheet_name='TopProducts', index=False)
        cat.to_excel(xw, sheet_name='Categories', index=False)
        city.to_excel(xw, sheet_name='Cities', index=False)
        pay.to_excel(xw, sheet_name='Payments', index=False)
        status.to_excel(xw, sheet_name='Status', index=False)
        if profit is not None:
            pd.DataFrame([profit]).to_excel(xw, sheet_name='Profit', index=False)
        if pprod is not None:
            pprod.to_excel(xw, sheet_name='ProfitProducts', index=False)
        if pcat is not None:
            pcat.to_excel(xw, sheet_name='ProfitCats', index=False)
        if basket is not None and len(basket):
            basket.head(50).to_excel(xw, sheet_name='Basket', index=False)
        if cohort is not None:
            cohort.to_excel(xw, sheet_name='Cohort')
        if forecast_info is not None:
            pd.DataFrame([forecast_info]).to_excel(xw, sheet_name='Forecast', index=False)
        if forecast_series is not None and len(forecast_series):
            forecast_series.to_excel(xw, sheet_name='ForecastSeries', index=False)
        if rfm is not None and len(rfm):
            rfm.to_excel(xw, sheet_name='RFM', index=False)
        if seg is not None and len(seg):
            seg.to_excel(xw, sheet_name='RFM_Segments', index=False)
        if restock is not None and len(restock):
            restock.to_excel(xw, sheet_name='Restock', index=False)
        if sentiment_summary is not None and len(sentiment_summary):
            sentiment_summary.to_excel(xw, sheet_name='Sentiment', index=False)
        if sentiment_byprod is not None and len(sentiment_byprod):
            sentiment_byprod.to_excel(xw, sheet_name='SentimentByProd', index=False)
        if alerts is not None and len(alerts):
            alerts.to_excel(xw, sheet_name='Alerts', index=False)
        if recs is not None and len(recs):
            recs.to_excel(xw, sheet_name='Recommend', index=False)
        if repeat is not None:
            pd.DataFrame([repeat]).to_excel(xw, sheet_name='Repeat', index=False)
        if quality is not None and len(quality):
            quality.to_excel(xw, sheet_name='Quality', index=False)
        if churn_summary is not None and len(churn_summary):
            churn_summary.to_excel(xw, sheet_name='Churn', index=False)
        if churn is not None and len(churn):
            churn.head(200).to_excel(xw, sheet_name='ChurnDetail', index=False)
        if tiers is not None and len(tiers):
            tiers.to_excel(xw, sheet_name='PriceTiers', index=False)
        if tier_corr is not None and len(tier_corr):
            tier_corr.to_excel(xw, sheet_name='PriceCorr', index=False)
        if gaps is not None and len(gaps):
            gaps.to_excel(xw, sheet_name='Gaps', index=False)
        if recoverable is not None and len(recoverable):
            recoverable.to_excel(xw, sheet_name='Recoverable', index=False)
        if upsell is not None:
            pd.DataFrame([upsell]).to_excel(xw, sheet_name='Upsell', index=False)
    # styling pass
    try:
        import openpyxl
        wb = openpyxl.load_workbook(out)
        fills = {'KPIs': '1F4E78', 'Monthly': '2E75B6', 'TopProducts': '548235', 'Profit': '375623',
                 'RFM': '7F6000', 'Restock': 'C55A11', 'Sentiment': '7030A0', 'Alerts': 'C00000', 'Recommend': '2E75B6'}
        for ws in wb.worksheets:
            _style_sheet(ws, fills.get(ws.title, '1F4E78'))
        wb.save(out)
    except Exception:
        pass
    return str(out)

